"""从文本类文档中提取纯文本，供知识库入口复用。

这里只处理已经是文字载体的文件；音视频转写属于直播采集链路，不在本模块内。
"""

from __future__ import annotations

import io
import zipfile
import zlib

from fastapi import HTTPException


TEXT_DOCUMENT_EXTENSIONS = (
    ".txt", ".md", ".csv", ".json", ".xlsx", ".xlsm", ".docx", ".pdf",
)

_PDF_MAX_FILE_BYTES = 30 * 1024 * 1024
_PDF_MAX_PAGES = 500
_PDF_MAX_CONTENT_STREAMS_PER_PAGE = 256
_PDF_MAX_PAGE_COMPRESSED_CONTENT = 8 * 1024 * 1024
_PDF_MAX_PAGE_DECODED_CONTENT = 16 * 1024 * 1024
_PDF_MAX_TOTAL_DECODED_CONTENT = 64 * 1024 * 1024


def _guard_office_zip(data: bytes, max_uncompressed: int = 64 * 1024 * 1024) -> None:
    """拒绝加密包和高膨胀 Office ZIP，避免小压缩包解压耗尽内存。"""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        members = archive.infolist()
        if any(member.flag_bits & 0x1 for member in members):
            raise ValueError("不支持加密的 Office 文档")
        total = sum(member.file_size for member in members)
        compressed = max(1, sum(member.compress_size for member in members))
        if total > max_uncompressed or total / compressed > 200:
            raise ValueError("文档解压后过大或压缩比异常")


def _bounded(text: str, max_chars: int | None) -> str:
    return text if max_chars is None else text[:max_chars + 1]


def _decode_plain_text(data: bytes) -> str:
    encodings = ["utf-8-sig", "utf-8", "gb18030", "latin-1"]
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):
        encodings.insert(0, "utf-16")
    for encoding in encodings:
        try:
            value = data.decode(encoding)
        except UnicodeDecodeError:
            continue
        controls = sum(
            1 for char in value
            if ord(char) < 32 and char not in {"\n", "\r", "\t", "\f"}
        )
        # latin-1 会“成功”解码任意二进制；文本性校验防止把改后缀的媒体送进 LLM。
        if "\x00" in value or (value and controls / len(value) > 0.02):
            continue
        return value
    raise ValueError("文件内容不像可解析的纯文本")


def _pdf_content_streams(page) -> list:
    """返回页面直接的 /Contents 流，不在这里解码。"""
    contents = page.get("/Contents")
    if contents is None:
        return []
    resolved = contents.get_object()
    streams = list(resolved) if isinstance(resolved, list) else [resolved]
    if len(streams) > _PDF_MAX_CONTENT_STREAMS_PER_PAGE:
        raise ValueError(
            f"PDF 单页内容流过多（最多 {_PDF_MAX_CONTENT_STREAMS_PER_PAGE} 个）"
        )
    result = []
    for stream in streams:
        stream = stream.get_object()
        if not hasattr(stream, "get_data"):
            raise ValueError("PDF 页面内容流格式异常")
        result.append(stream)
    return result


def _pdf_filter_names(stream) -> tuple[str, ...]:
    filters = stream.get("/Filter")
    if filters is None:
        return ()
    filters = filters.get_object()
    values = filters if isinstance(filters, list) else [filters]
    return tuple(str(value.get_object()) for value in values)


def _guard_flate_expansion(raw: bytes, decoded_budget: int) -> None:
    """对常见 Flate 内容流做有界解压预检，不物化完整高膨胀结果。"""
    for window_bits in (zlib.MAX_WBITS, -zlib.MAX_WBITS):
        decompressor = zlib.decompressobj(window_bits)
        try:
            output = decompressor.decompress(raw, decoded_budget + 1)
        except zlib.error:
            continue
        if len(output) > decoded_budget or decompressor.unconsumed_tail:
            raise ValueError("PDF 单页解压后内容流过大")
        return
    # 容错留给 pypdf：它对部分非标准 Flate 有额外恢复路径。


def _guard_pdf_page_content(page, total_decoded: int) -> int:
    """在 extract_text 前约束页面直接内容流的数量和体积。"""
    streams = _pdf_content_streams(page)
    raw_sizes = [len(getattr(stream, "_data", b"") or b"") for stream in streams]
    if sum(raw_sizes) > _PDF_MAX_PAGE_COMPRESSED_CONTENT:
        raise ValueError("PDF 单页压缩内容流过大")

    page_decoded = 0
    for stream in streams:
        remaining_page = _PDF_MAX_PAGE_DECODED_CONTENT - page_decoded
        remaining_total = _PDF_MAX_TOTAL_DECODED_CONTENT - total_decoded - page_decoded
        budget = min(remaining_page, remaining_total)
        if budget < 0:
            raise ValueError("PDF 解压后内容流总量过大")
        raw = getattr(stream, "_data", b"") or b""
        filters = _pdf_filter_names(stream)
        if filters and filters[0] in ("/FlateDecode", "/Fl"):
            _guard_flate_expansion(raw, budget)
        decoded = stream.get_data()
        decoded_size = len(decoded)
        if decoded_size > budget:
            if decoded_size + page_decoded > _PDF_MAX_PAGE_DECODED_CONTENT:
                raise ValueError("PDF 单页解压后内容流过大")
            raise ValueError("PDF 解压后内容流总量过大")
        page_decoded += decoded_size
    return total_decoded + page_decoded


def extract_document_text(filename: str, data: bytes, max_chars: int | None = 2_000_000) -> str:
    """提取支持文档中的可检索文字；不支持的类型返回明确的 400。"""
    name = (filename or "").lower()
    if name.endswith((".txt", ".md", ".csv", ".json")):
        return _bounded(_decode_plain_text(data), max_chars)

    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        _guard_office_zip(data)
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        rows: list[str] = []
        length = 0
        try:
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    cells = [str(cell) for cell in row if cell is not None]
                    if cells:
                        line = " ".join(cells)
                        rows.append(line)
                        length += len(line) + 1
                        if max_chars is not None and length > max_chars:
                            return _bounded("\n".join(rows), max_chars)
        finally:
            wb.close()
        return "\n".join(rows)

    if name.endswith(".docx"):
        import docx

        _guard_office_zip(data)
        document = docx.Document(io.BytesIO(data))
        blocks = [p.text for p in document.paragraphs if p.text.strip()]
        # 课程讲义偶尔把逐字稿放在表格中，不能只读取段落。
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    blocks.append(" | ".join(cells))
        return _bounded("\n".join(blocks), max_chars)

    if name.endswith(".pdf"):
        from pypdf import PdfReader

        if len(data) > _PDF_MAX_FILE_BYTES:
            raise ValueError("PDF 文件过大（最大 30MB）")
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            raise ValueError("不支持加密 PDF")
        if len(reader.pages) > _PDF_MAX_PAGES:
            raise ValueError(f"PDF 页数过多（最多 {_PDF_MAX_PAGES} 页）")
        pages: list[str] = []
        length = 0
        total_decoded = 0
        for page in reader.pages:
            # 先检查直接页面内容流，再调用可能建立操作列表的 extract_text。
            # 这是资源边界而非完整沙箱：Form XObject 递归/重复引用仍可能放大工作量。
            total_decoded = _guard_pdf_page_content(page, total_decoded)
            page_text = page.extract_text() or ""
            pages.append(page_text)
            length += len(page_text) + 1
            if max_chars is not None and length > max_chars:
                return _bounded("\n".join(pages), max_chars)
        return "\n".join(pages)

    raise HTTPException(400, f"不支持的文件类型：{filename}")
