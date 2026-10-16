"""复盘场次接口：建场/解析/起停/主播切换/查询(可按主播)/下载/彻底删除/视频代理 + WS。"""

from __future__ import annotations

import asyncio
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
from datetime import datetime, timedelta
from pathlib import Path
from types import SimpleNamespace
from urllib.parse import quote

from fastapi import (
    APIRouter, Body, Depends, HTTPException, Query, Request, WebSocket, WebSocketDisconnect,
)
from fastapi.responses import (
    FileResponse, PlainTextResponse, RedirectResponse, Response, StreamingResponse,
)
from starlette.concurrency import run_in_threadpool
from starlette.background import BackgroundTask
from ..core import media_store
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from ..bus import bus
from ..config import get_settings, public_url
from ..core import douyin
from ..core.embed import retrieve_detailed
from ..core.live import open_live_stream
from ..core.live_link import normalize_live_link
from ..core.asr_stream import ASRConfigurationError, ensure_real_asr_configured
from ..core.llm import get_llm
from ..core.pipeline import get_running, start_pipeline, stop_pipeline
from ..core.recovery import recover_session
from ..core.safe_error import safe_error_text
from ..db import get_db
from ..models import (
    AudioStatus,
    CompassPoint,
    Critique,
    CritiqueLevel,
    LiveRoom,
    OcrRecord,
    ReplaySession,
    SessionType,
    SessionRhythmAnalysis,
    SessionStatus,
    Streamer,
    StreamerTurn,
    Transcript,
)
from ..schemas import (
    AudioOut,
    CritiqueOut,
    SessionCreate,
    SessionOut,
    StreamerOut,
    StreamerReviewOut,
    SummaryRegenerateIn,
    SwitchIn,
    TranscriptOut,
)
from ..storage.audio import get_audio, hard_delete_audio, find_part_file

router = APIRouter(prefix="/api/sessions", tags=["sessions"])
_ACTIVE_SESSION_STATUSES = (
    SessionStatus.created,
    SessionStatus.parsing,
    SessionStatus.live,
    SessionStatus.processing,
)


def _json_dict(value: str | None) -> dict:
    try:
        parsed = json.loads(value or "{}")
    except (json.JSONDecodeError, TypeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


@router.post("/{session_id}/summary/regenerate")
async def regenerate_summary(session_id: int, body: SummaryRegenerateIn,
                             db: Session = Depends(get_db)):
    """Regenerate the authoritative summary without damaging the previous report on failure."""
    row = db.get(ReplaySession, session_id)
    if row is None:
        raise HTTPException(404, "场次不存在")
    if row.session_type == SessionType.observation:
        raise HTTPException(409, detail={
            "code": "OBSERVATION_SUMMARY_NOT_SUPPORTED",
            "message": "观摩学习场次不生成自家直播整场复盘",
        })
    if row.ended_at is None or row.status in _ACTIVE_SESSION_STATUSES:
        raise HTTPException(409, "仅已结束场次可以重新生成整场总结")
    try:
        from ..core.summary_service import generate_summary
        data = await generate_summary(session_id, body.streamer_id)
        return {"ok": True, "streamer_id": body.streamer_id, "summary": data}
    except ValueError as exc:
        raise HTTPException(422, safe_error_text(exc)) from None
    except Exception as exc:  # noqa: BLE001
        from ..core.llm import LLMCallError
        if isinstance(exc, LLMCallError):
            code_map = {
                "LLM_REQUEST_TIMEOUT": (504, "SUMMARY_REQUEST_TIMEOUT", "整场总结请求超时，请稍后重试"),
                "LLM_NETWORK_ERROR": (503, "SUMMARY_NETWORK_ERROR", "模型网络连接中断，请稍后重试"),
                "LLM_RESPONSE_INCOMPLETE": (503, "SUMMARY_RESPONSE_INCOMPLETE", "模型返回内容不完整，请稍后重试"),
                "LLM_REQUEST_REJECTED": (502, "SUMMARY_REQUEST_REJECTED", "模型拒绝了总结请求，请检查密钥、权限、余额或参数配置"),
                "LLM_RATE_LIMITED": (503, "SUMMARY_RATE_LIMITED", "模型服务触发限流，请稍后重试"),
                "LLM_PROVIDER_UNAVAILABLE": (503, "SUMMARY_PROVIDER_UNAVAILABLE", "模型服务暂不可用，请稍后重试"),
            }
            status, code, message = code_map.get(
                exc.code, (503, "SUMMARY_GENERATION_FAILED", "整场总结生成失败，请稍后重试"),
            )
            raise HTTPException(status, detail={"code": code, "message": message}) from None
        raise HTTPException(503, detail={
            "code": "SUMMARY_GENERATION_FAILED",
            "message": "整场总结生成失败，请稍后重试",
        }) from None


def _rhythm_row_dict(analysis: SessionRhythmAnalysis) -> dict:
    return {
        "available": True,
        "id": analysis.id,
        "session_id": analysis.session_id,
        "streamer_id": analysis.streamer_id,
        "template_id": analysis.template_id,
        "selection_mode": analysis.selection_mode,
        "declared_execution": analysis.declared_execution,
        "declared_at": (
            analysis.declared_at.isoformat(timespec="seconds")
            if analysis.declared_at else None
        ),
        "declaration_start_ms": analysis.declaration_start_ms,
        "status": analysis.status,
        "template": _json_dict(analysis.template_snapshot),
        "actual": _json_dict(analysis.extracted),
        "comparison": _json_dict(analysis.comparison),
        "validation": _json_dict(analysis.validation),
        "match_score": analysis.match_score,
        "effect_score": analysis.effect_score,
        "effective": analysis.effective,
        "error_msg": analysis.error_msg,
    }


def _projected_summary_data(raw_data: str | None, analysis: SessionRhythmAnalysis) -> str:
    from ..core.review.rhythm_analysis import project_report_data

    data = _json_dict(raw_data)
    project_report_data(data, _rhythm_row_dict(analysis))
    return json.dumps(data, ensure_ascii=False)


def _session_out(db: Session, row: ReplaySession) -> SessionOut:
    try:
        lineup_ids = json.loads(row.lineup) if row.lineup else []
    except json.JSONDecodeError:
        lineup_ids = []
    smap = {s.id: s for s in db.query(Streamer).filter(Streamer.id.in_(lineup_ids)).all()} \
        if lineup_ids else {}
    lineup = [StreamerOut.model_validate(smap[i]) for i in lineup_ids if i in smap]
    cur = db.get(Streamer, row.current_streamer_id) if row.current_streamer_id else None
    from ..core.compliance_rules import get_session_snapshot
    return SessionOut(
        id=row.id, title=row.title, share_link=row.share_link, room_id=row.room_id,
        stream_url=row.stream_url, status=row.status.value, error_code=row.error_code,
        report_status=getattr(row, "report_status", "pending") or "pending",
        deep_status=getattr(row, "deep_status", "not_requested") or "not_requested",
        error_msg=safe_error_text(row.error_msg, fallback="", limit=180),
        started_at=row.started_at, ended_at=row.ended_at, created_at=row.created_at,
        audio=AudioOut.model_validate(row.audio) if row.audio else None,
        current_streamer_id=row.current_streamer_id,
        current_streamer_name=cur.name if cur else "",
        lineup=lineup,
        session_type=getattr(row.session_type, "value", row.session_type or "normal"),
        compliance_snapshot=get_session_snapshot(row.id, db=db),
    )


def _active_session_for_link(db: Session, normalized_link: str) -> ReplaySession | None:
    if not normalized_link:
        return None
    return db.execute(select(ReplaySession).where(
        ReplaySession.normalized_share_link == normalized_link,
        ReplaySession.status.in_(_ACTIVE_SESSION_STATUSES),
        ReplaySession.ended_at.is_(None),
    ).order_by(ReplaySession.id)).scalars().first()


def _duplicate_session_conflict(db: Session, row: ReplaySession) -> HTTPException:
    return HTTPException(status_code=409, detail={
        "code": "ACTIVE_SESSION_EXISTS",
        "message": "同一直播链接已有未结束场次，请进入原场次",
        # The client only needs identity and display state to enter the
        # existing session. Never echo signed CDN pull URLs in an error body.
        "session": {
            "id": row.id,
            "title": row.title,
            "status": row.status.value,
            "started_at": row.started_at.isoformat() if row.started_at else None,
        },
    })


def _automatic_session_title(
    native_title: str = "", anchor_nickname: str = "", *, now: datetime | None = None
) -> str:
    """原生标题优先；无标题时严格按“主播昵称+时间/直播场次+时间”兜底。"""
    moment = now or datetime.now()
    title = (native_title or "").strip()
    nickname = (anchor_nickname or "").strip()
    if title:
        base = f"{nickname}｜{title}" if nickname and nickname not in title else title
    elif nickname:
        base = nickname
    else:
        base = "直播场次"
    return f"{base[:100]} · {moment:%m-%d %H:%M}"


@router.post("", response_model=SessionOut)
async def create_session(payload: SessionCreate, db: Session = Depends(get_db)):
    normalized_link = normalize_live_link(payload.share_link or payload.manual_stream_url)
    existing = _active_session_for_link(db, normalized_link)
    if existing is not None:
        raise _duplicate_session_conflict(db, existing)
    is_observation = payload.session_type == "observation"
    # 服务端同样执行就绪门禁，防止绕过前端直接创建不可运行的场次。
    from ..main import ready as system_ready
    # Chromium 使用 Playwright 同步 API；异步建场处理器必须在线程池中执行检查，
    # 否则 Playwright 会因当前线程已有 asyncio 事件循环而误报“无法启动”。
    readiness = await run_in_threadpool(system_ready)
    blockers = readiness["blockers"]
    if is_observation:
        # 观摩场次不走实时 LLM，不能仅因模型未配置而阻止录音、逐字稿和弹幕采集。
        blockers = [
            f'{item["label"]}：{item["message"]}'
            for name, item in readiness.get("components", {}).items()
            if name != "llm" and not item["ok"]
        ]
    if blockers:
        raise HTTPException(503, "系统未就绪，无法新建场次：" + "；".join(blockers))
    requested_ids = list(dict.fromkeys(payload.streamer_ids))
    invalid_ids = [i for i in requested_ids if db.get(Streamer, i) is None]
    if invalid_ids:
        raise HTTPException(422, f"主播不存在：{', '.join(map(str, invalid_ids))}")
    if (not is_observation and db.query(Streamer).filter(Streamer.deleted_at.is_(None)).count()
            and not requested_ids and not payload.allow_unassigned):
        raise HTTPException(422, "请选择本场主播；如确需无主播场次，请明确选择“暂不指定主播”")
    lineup = [] if is_observation else requested_ids
    row = ReplaySession(
        title=payload.title,
        share_link=normalize_live_link(payload.share_link) or payload.share_link.strip(),
        normalized_share_link=normalized_link,
        lineup=json.dumps(lineup),
        current_streamer_id=(lineup[0] if lineup else None),
        session_type=SessionType.observation if is_observation else SessionType.normal,
        compliance_rule_ids=json.dumps(list(dict.fromkeys(payload.compliance_rule_ids))),
        status=SessionStatus.parsing,
    )
    db.add(row)
    try:
        db.commit()
    except IntegrityError:
        # 两个请求同时通过上方只读检查时，由数据库部分唯一索引裁决。
        db.rollback()
        existing = _active_session_for_link(db, normalized_link)
        if existing is not None:
            raise _duplicate_session_conflict(db, existing) from None
        raise
    db.refresh(row)

    room_title = ""
    anchor_nickname = ""
    try:
        res = await douyin.resolve_stream(
            share_link=payload.share_link, manual_url=payload.manual_stream_url,
            cookie=get_settings().douyin_cookie,
        )
        row.room_id = res.room_id
        row.stream_url = res.stream_url
        row.status = SessionStatus.created
        room_title = (res.title or "").strip()
        anchor_nickname = (res.anchor_nickname or "").strip()
    except douyin.ParseError:
        row.status = SessionStatus.created
        row.error_code = "STREAM_PARSE_FAILED"
        row.error_msg = "直播源解析失败，请确认链接有效或填写可控拉流地址"
    # 自动命名：仅信任直播结构化数据中的主播/原生标题，不读取广告/SEO 页面标题。
    if not (payload.title or "").strip():
        if not anchor_nickname and lineup:
            assigned = db.get(Streamer, lineup[0])
            anchor_nickname = assigned.name if assigned else ""
        row.title = _automatic_session_title(room_title, anchor_nickname)
    # 收录/更新「常用直播间」列表（按 room_id 去重），之后可从列表复选
    if row.room_id:
        rm = db.query(LiveRoom).filter(
            LiveRoom.room_id == row.room_id, LiveRoom.deleted_at.is_(None),
        ).first()
        if not rm:
            rm = LiveRoom(room_id=row.room_id)
            db.add(rm)
        if room_title or anchor_nickname:
            rm.title = room_title or anchor_nickname
        if payload.share_link:
            rm.share_link = payload.share_link
        # 手动建场解析成功也学习账号级恒定 ID（场控救火粘链接的场景）
        if getattr(res, "sec_uid", "") and not (getattr(rm, "sec_user_id", "") or ""):
            rm.sec_user_id = res.sec_uid
        rm.last_used_at = datetime.now()
    db.commit()
    db.refresh(row)
    return _session_out(db, row)


@router.get("", response_model=list[SessionOut])
def list_sessions(include_error: bool = Query(False), db: Session = Depends(get_db)):
    # 轮询试探失败等 error 场次默认隐藏（对用户是噪音）；排障时可传 include_error=1
    q = select(ReplaySession).order_by(ReplaySession.id.desc())
    if not include_error:
        q = q.where(ReplaySession.status != SessionStatus.error)
    rows = db.execute(q).scalars().all()
    return [_session_out(db, r) for r in rows]


@router.get("/{session_id}", response_model=SessionOut)
def get_session(session_id: int, db: Session = Depends(get_db)):
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    return _session_out(db, row)


@router.post("/{session_id}/start", response_model=SessionOut)
async def start_session(session_id: int,
                        streamer_ids: list[int] | None = Body(None, embed=True),
                        db: Session = Depends(get_db),
                        compliance_rule_ids: list[int] | None = Body(None, embed=True)):
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    # 防重复开播污染历史：仅草稿(created)或上次失败(error)可开播；done/processing/live 拒绝
    if row.status not in (SessionStatus.created, SessionStatus.error):
        raise HTTPException(409, f"场次当前为「{row.status.value}」，不能重复开播；如需重跑请新建场次")
    # 开播时应用当前排班：草稿是粘链接时自动建的、那一刻往往还没选主播，
    # 故以开播一刻前端传来的排班为准，避免主播丢失（仍显示"未指定主播"）。
    is_observation = row.session_type == SessionType.observation
    if streamer_ids is not None and not is_observation:
        lineup = [i for i in streamer_ids if db.get(Streamer, i)]
        row.lineup = json.dumps(lineup)
        row.current_streamer_id = lineup[0] if lineup else None
        db.commit()
        db.refresh(row)
    elif is_observation and (row.lineup or row.current_streamer_id is not None):
        row.lineup = "[]"
        row.current_streamer_id = None
        db.commit()
        db.refresh(row)
    if row.stream_url:
        try:
            ensure_real_asr_configured()
        except ASRConfigurationError as exc:
            raise HTTPException(status_code=503, detail={
                "code": "ASR_NOT_READY",
                "message": str(exc),
            }) from None
    if compliance_rule_ids is None:
        try:
            compliance_rule_ids = json.loads(row.compliance_rule_ids or "[]")
        except (json.JSONDecodeError, TypeError):
            compliance_rule_ids = []
    else:
        row.compliance_rule_ids = json.dumps(list(dict.fromkeys(compliance_rule_ids)))
        db.commit()
    try:
        from ..core.compliance_rules import lock_session_snapshot
        lock_session_snapshot(session_id, compliance_rule_ids)
    except (LookupError, ValueError) as exc:
        raise HTTPException(422, detail={
            "code": "COMPLIANCE_SNAPSHOT_INVALID", "message": str(exc),
        }) from None
    # 无拉流地址时走合成测试流（mock 识别）；有地址则真实拉流 + 真实 ASR/OCR
    start_pipeline(session_id)
    return _session_out(db, row)


@router.post("/{session_id}/stop", response_model=SessionOut)
async def stop_session(session_id: int, db: Session = Depends(get_db)):
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    if row.status in (SessionStatus.done, SessionStatus.error):
        if row.status == SessionStatus.done and row.session_type == SessionType.observation:
            from ..core.observation import compute_and_store_heat
            await asyncio.to_thread(compute_and_store_heat, session_id)
        return _session_out(db, row)
    running = stop_pipeline(session_id)
    if not running and row.status in (SessionStatus.live, SessionStatus.processing):
        # The worker disappeared (typically a restart).  Finish the persisted
        # artifacts once; recovery never invokes ASR/LLM or generates a report.
        await recover_session(session_id)
        db.expire_all()
        row = db.get(ReplaySession, session_id)
        if row and row.status == SessionStatus.done and row.session_type == SessionType.observation:
            from ..core.observation import compute_and_store_heat
            await asyncio.to_thread(compute_and_store_heat, session_id)
    return _session_out(db, row)


@router.post("/{session_id}/switch", response_model=SessionOut)
async def switch_streamer(session_id: int, payload: SwitchIn, db: Session = Depends(get_db)):
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    if row.session_type == SessionType.observation:
        raise HTTPException(409, "观摩学习场次不绑定或切换自家主播")
    try:
        lineup = json.loads(row.lineup) if row.lineup else []
    except json.JSONDecodeError:
        lineup = []

    target = payload.streamer_id
    if payload.next:
        if not lineup:
            raise HTTPException(400, "未配置排班，无法自动切下一位")
        cur = row.current_streamer_id
        idx = lineup.index(cur) if cur in lineup else -1
        target = lineup[(idx + 1) % len(lineup)]
    if target is None or not db.get(Streamer, target):
        raise HTTPException(400, "目标主播无效")

    p = get_running(session_id)
    if p:
        await p.switch_streamer(target)   # 直播中：记班次 + 推送
        db.refresh(row)
    else:
        row.current_streamer_id = target  # 未开播：仅设定开播主播
        db.commit()
    return _session_out(db, row)


@router.get("/{session_id}/streamers", response_model=list[StreamerReviewOut])
def session_streamers(session_id: int, db: Session = Depends(get_db)):
    """按主播分组的复盘聚合（每个主播的句数、告警数、班次、整场总结）。"""
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    # 哪些主播正在生成 v2：必须先于下面读数据拍下来。反过来的话，任务若恰好在两步之间存完收尾，
    # 页面会看到「没在跑 + 结果没变」而误判成这次没成。
    from ..core.review.v2_jobs import running_jobs
    v2_running = running_jobs(session_id)

    turns = db.execute(
        select(StreamerTurn).where(StreamerTurn.session_id == session_id)
        .order_by(StreamerTurn.seq)
    ).scalars().all()
    transcripts = db.execute(
        select(Transcript).where(Transcript.session_id == session_id)
    ).scalars().all()
    critiques = db.execute(
        select(Critique).where(Critique.session_id == session_id)
    ).scalars().all()
    rhythm_rows = db.execute(select(SessionRhythmAnalysis).where(
        SessionRhythmAnalysis.session_id == session_id,
    )).scalars().all()
    rhythm_by_streamer = {analysis.streamer_id: analysis for analysis in rhythm_rows}

    try:
        lineup = json.loads(row.lineup) if row.lineup else []
    except json.JSONDecodeError:
        lineup = []
    ordered: list[int | None] = list(lineup)
    for coll in (turns, transcripts, critiques, rhythm_rows):
        for x in coll:
            if x.streamer_id not in ordered:
                ordered.append(x.streamer_id)

    out: list[StreamerReviewOut] = []
    for sid in ordered:
        s = db.get(Streamer, sid) if sid else None
        summary = next((c for c in critiques
                        if c.streamer_id == sid
                        and c.level == CritiqueLevel.summary
                        and c.category == "整场总结"), None)
        analysis = rhythm_by_streamer.get(sid)
        summary_out = CritiqueOut.model_validate(summary) if summary else None
        if summary_out is not None and analysis is not None:
            summary_out.data = _projected_summary_data(summary.data, analysis)
        elif analysis is not None:
            summary_out = CritiqueOut(
                id=-analysis.id,
                level="summary",
                severity="info",
                category="整场总结",
                start_ms=0,
                end_ms=0,
                quote="",
                comment="R3 话术节奏分析已生成",
                score=None,
                data=_projected_summary_data(None, analysis),
                streamer_id=sid,
            )
        out.append(StreamerReviewOut(
            streamer_id=sid or 0,
            name=s.name if s else "未指定主播",
            color=s.color if s else "",
            sentence_count=sum(1 for t in transcripts if t.streamer_id == sid),
            flag_count=sum(1 for c in critiques if c.streamer_id == sid
                           and c.level in (CritiqueLevel.sentence, CritiqueLevel.screen)),
            turns=[{"start_ms": t.start_ms, "end_ms": t.end_ms}
                   for t in turns if t.streamer_id == sid],
            summary=summary_out,
            overall_score=summary.score if summary else None,
            v2_job=v2_running.get(sid or None),
        ))
    return out


@router.get("/{session_id}/transcripts", response_model=list[TranscriptOut])
def list_transcripts(session_id: int, streamer_id: int | None = None,
                     db: Session = Depends(get_db)):
    stmt = select(Transcript).where(Transcript.session_id == session_id)
    if streamer_id is not None:
        stmt = stmt.where(Transcript.streamer_id == streamer_id)
    return db.execute(stmt.order_by(Transcript.seq)).scalars().all()


@router.get("/{session_id}/roles")
def session_roles(session_id: int, streamer_id: int | None = None,
                  db: Session = Depends(get_db)):
    """主播 / 场控双线分工（从内容自动分搭话）。场控当前不出声则 assist_count=0。"""
    from ..core.roles import analyze_roles
    stmt = select(Transcript).where(Transcript.session_id == session_id)
    if streamer_id is not None:
        stmt = stmt.where(Transcript.streamer_id == streamer_id)
    trs = db.execute(stmt.order_by(Transcript.seq)).scalars().all()
    return analyze_roles(trs)


@router.post("/{session_id}/emotion/recompute")
def recompute_emotion(session_id: int, streamer_id: int | None = None,
                      db: Session = Depends(get_db)):
    """用 qwen3-asr-flash 对每个点评段识别情绪(7类)，写回段级 + 聚合整场分布。"""
    from ..core import qwen_emotion
    from ..core.emotion_recompute import recompute_session_emotions
    if not qwen_emotion.enabled():
        raise HTTPException(400, "未配置百炼 key，无法用 Qwen 识别情绪")
    if not db.get(ReplaySession, session_id):
        raise HTTPException(404, "场次不存在")
    res = recompute_session_emotions(session_id, streamer_id)
    if res.get("error") == "no_audio":
        raise HTTPException(404, "本场无录音，无法识别情绪")
    if res.get("error") == "no_source":
        raise HTTPException(404, "源音频不可用（本地已清理且无 COS 备份）")
    from ..core.emotion_context import recompute_session_contexts
    res["emotion_context"] = recompute_session_contexts(session_id, streamer_id)
    return res


_COMPASS_JS = r'''javascript:(function(){var U="__URL__",T="__TOKEN__";function g(){var m={},KS=["在线","观看","累计观看","成交金额","成交","GMV","转化率","转化","点赞","曝光","进入率","停留","客单价","订单","销量","粉丝","涨粉"];var els=document.querySelectorAll("body *");for(var i=0;i<els.length;i++){var e=els[i];if(e.children.length)continue;var t=(e.textContent||"").trim();if(!t||t.length>10)continue;for(var k=0;k<KS.length;k++){if(t.indexOf(KS[k])>=0){var v="",s=e.nextElementSibling;if(s)v=(s.textContent||"").trim();if(!v||!/[0-9]/.test(v)){var p=e.parentElement;if(p)v=(p.textContent||"").replace(t,"").trim()}if(v&&/[0-9]/.test(v))m[KS[k]]=v.slice(0,20)}}}try{navigator.sendBeacon(U,JSON.stringify({token:T,data:m}))}catch(x){}}g();if(window.__cmp)clearInterval(window.__cmp);window.__cmp=setInterval(g,30000);alert("罗盘上报已启动(每30秒),保持本页打开。停止请刷新本页。")})();'''


def _build_bookmarklet(url: str, token: str) -> str:
    return _COMPASS_JS.replace("__URL__", url).replace("__TOKEN__", token)


@router.get("/{session_id}/compass/setup")
def compass_setup(session_id: int, request: Request, db: Session = Depends(get_db)):
    """生成/返回罗盘上报书签（bookmarklet）+ token。"""
    import secrets
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    if not row.compass_token:
        row.compass_token = secrets.token_urlsafe(18)
        db.commit()
    # 对外地址取 PUBLIC_BASE_URL（含 /fupan 前缀），不再按 Host 拼——经 workspace.example.test/fupan 访问时 Host 不带前缀
    url = public_url(f"/api/sessions/{session_id}/compass")
    return {"token": row.compass_token, "url": url,
            "bookmarklet": _build_bookmarklet(url, row.compass_token)}


@router.post("/{session_id}/compass")
async def compass_report(session_id: int, request: Request, db: Session = Depends(get_db)):
    """罗盘书签上报数据点（公开端点，靠 token 校验；不走 cookie 鉴权）。"""
    import hmac
    row = db.get(ReplaySession, session_id)
    if not row or not row.compass_token:
        raise HTTPException(404, "场次不存在或未启用罗盘")
    if row.ended_at:                          # 收播后不再收，避免晚到点偏移错乱
        raise HTTPException(409, "场次已结束，无需再上报")
    body = await request.body()
    if len(body) > 8192:                      # 单次体积上限，防公开端点被灌爆
        raise HTTPException(413, "数据过大")
    try:
        payload = json.loads(body)
    except (json.JSONDecodeError, ValueError):
        raise HTTPException(400, "数据格式错误")
    if not hmac.compare_digest(str(payload.get("token") or ""), row.compass_token):
        raise HTTPException(403, "token 无效")
    if db.query(CompassPoint).filter(CompassPoint.session_id == session_id).count() >= 5000:
        raise HTTPException(429, "本场罗盘数据点已达上限")
    data = payload.get("data") or {}
    ds = json.dumps(data, ensure_ascii=False)[:4000]   # data 长度上限
    start_ms = max(0, int((datetime.now() - row.started_at).total_seconds() * 1000)) if row.started_at else 0
    db.add(CompassPoint(session_id=session_id, start_ms=start_ms, data=ds))
    db.commit()
    return {"ok": True}


@router.get("/{session_id}/compass")
def compass_list(session_id: int, db: Session = Depends(get_db)):
    pts = db.execute(select(CompassPoint).where(CompassPoint.session_id == session_id)
                     .order_by(CompassPoint.start_ms)).scalars().all()
    def _d(p):
        try:
            return json.loads(p.data) if p.data else {}
        except json.JSONDecodeError:
            return {}
    return [{"start_ms": p.start_ms, "data": _d(p),
             "at": p.created_at.isoformat() if p.created_at else None} for p in pts]


@router.get("/{session_id}/danmaku_signals")
def danmaku_signals(session_id: int, db: Session = Depends(get_db)):
    """公屏弱信号分布 + 每类原文样本 + 归因提示；带上最近一次 AI 意向精析结果。"""
    from ..core.danmaku_refine import load_refine
    from ..core.danmaku_signal import analyze_session_danmaku
    out = analyze_session_danmaku(session_id)
    out["refine"] = load_refine(session_id)
    return out


@router.post("/{session_id}/danmaku_signals/refine")
async def danmaku_signals_refine(session_id: int, db: Session = Depends(get_db)):
    """AI 结合上下文精析购买意向（真实意向/疑似玩笑或跟风/咨询了解），结果落库。"""
    if not db.get(ReplaySession, session_id):
        raise HTTPException(404, "场次不存在")
    from ..core.danmaku_refine import refine_session_intents
    return await refine_session_intents(session_id)


@router.get("/{session_id}/attribution")
def attribution(session_id: int, db: Session = Depends(get_db)):
    """成交归因：成交时刻映射到话术段，判断话术驱动 vs 自然成交。"""
    from ..core.attribution import analyze_attribution
    return analyze_attribution(session_id)


@router.post("/{session_id}/reanalyze_v2")
async def reanalyze_v2(session_id: int, streamer_id: int | None = None,
                       only_failed: bool = False, db: Session = Depends(get_db)):
    """v2 深度复盘重跑（2026-07-15 改版）：对已完结场次生成五块分析并合入 summary.data.v2。
    不传 streamer_id 时对本场所有已有整场总结的主播都生成。幂等可重跑。
    同一（场次, 主播）已有生成在跑时不另起一份，本请求等那一份的结果（2026-09-30 防重复）。
    only_failed=1 只补跑上次失败/缺失的块，其余保留。请求一直挂到生成完才返回。"""
    from ..core.review import v2_jobs
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    rows = db.execute(select(Critique).where(
        Critique.session_id == session_id, Critique.level == CritiqueLevel.summary,
        Critique.category == "整场总结")).scalars().all()
    sids = list(dict.fromkeys(
        r.streamer_id for r in rows if streamer_id is None or r.streamer_id == streamer_id))
    db.close()      # 下面要等几分钟，不一直占着数据库连接
    if not sids:
        return {"ok": False, "error": "该场次没有整场总结（先完成收尾或指定正确主播）"}
    # 先把各主播的任务都登记上再等：等的过程中别处再点同一位主播，也只会等到这一份
    jobs = [(sid, *v2_jobs.start_job(session_id, sid, only_failed=only_failed)) for sid in sids]
    results = []
    for sid, job, created in jobs:
        outcome = await v2_jobs.wait_job(job)
        item = {"streamer_id": sid, "ok": outcome.ok,
                "blocks": [k for k in outcome.v2 if k != "generated_at"],
                "rerun": outcome.reran, "failed_blocks": outcome.failed,
                "joined": not created}
        if outcome.error:
            item["error"] = outcome.error
        results.append(item)
    bad = [item for item in results if not item["ok"]]
    if bad:
        return {"ok": False, "error": bad[0].get("error") or "生成失败", "results": results}
    return {"ok": True, "results": results}


@router.get("/{session_id}/critiques", response_model=list[CritiqueOut])
def list_critiques(session_id: int, level: str | None = None,
                   streamer_id: int | None = None, db: Session = Depends(get_db)):
    stmt = select(Critique).where(Critique.session_id == session_id)
    if level:
        stmt = stmt.where(Critique.level == CritiqueLevel(level))
    if streamer_id is not None:
        stmt = stmt.where(Critique.streamer_id == streamer_id)
    return db.execute(stmt.order_by(Critique.id)).scalars().all()


@router.get("/{session_id}/live.flv")
async def live_video(
    session_id: int,
    quality: str = Query("clear", pattern="^(smooth|clear|high|source)$"),
    db: Session = Depends(get_db),
):
    """实时直播画面（内存转发，不落盘）。默认清晰预览，可切换流畅/原画。"""
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    # 生产环境无拉流地址（未开播/链接无效/已结束）→ 明确报错，不回退 mock 彩条误导用户
    if not row.stream_url and get_settings().asr_provider != "mock":
        raise HTTPException(409, "该场次无实时画面：直播未开播、链接无效或场次已结束")
    return StreamingResponse(open_live_stream(row.stream_url, quality), media_type="video/x-flv")


@router.get("/{session_id}/audio/download")
def download_audio(session_id: int, streamer_id: int | None = None,
                   db: Session = Depends(get_db)):
    af = get_audio(db, session_id)
    if not af or af.status == AudioStatus.deleted:
        raise HTTPException(404, "媒体不存在或已删除")
    if af.status == AudioStatus.error:
        raise HTTPException(409, "媒体收尾失败，当前仅保留可恢复分片，不能作为成品下载")
    if streamer_id:   # 按主播出镜时段裁剪
        return _clip_audio_by_streamer(db, session_id, af, streamer_id)
    row = db.get(ReplaySession, session_id)
    kind = getattr(af, "media_kind", "audio")
    fname = _media_filename(db, row, None, kind) if row else af.filename
    if af.cos_key:   # 已上 COS：跳转临时签名 URL（附可读下载名）
        if af.cos_key.endswith(".m3u8"):
            # HLS 场次：下载给伴随音频（完整视频请在线回看，切片体积大不提供整包下载）
            akey = media_store.audio_track_key(af, session_id)
            if akey:
                aname = fname.rsplit(".", 1)[0] + ".m4a" if "." in fname else f"session_{session_id}.m4a"
                url = media_store.presigned_url(akey, expires=7200, filename=aname)
                if url:
                    return RedirectResponse(url)
            raise HTTPException(409, "本场为切片直传录制，暂无可下载的整段文件（可在线回看）")
        url = media_store.presigned_url(af.cos_key, expires=7200, filename=fname)
        if url:
            return RedirectResponse(url)
    p = Path(af.path)
    if not p.exists():
        # 录制中成品未合并（只有 .partN 分段）→ 给明确原因，而不是笼统的"文件缺失"
        if row and row.status in (SessionStatus.live, SessionStatus.processing) and find_part_file(af.path):
            raise HTTPException(409, "本场直播还在录制中，下播并完成处理后即可导出完整音视频")
        raise HTTPException(404, "磁盘文件缺失（可能已被删除）")
    mt = "video/mp4" if kind == "video" else "audio/mp4"
    return FileResponse(path=str(p), filename=fname, media_type=mt)


def _audio_source(af) -> str | None:
    """裁剪用源：本地文件优先，否则 COS 预签名 URL（ffmpeg 可直接读，支持 Range）；
    HLS 场次用伴随音频小文件（session_N.m4a）；录制中回退最新分段文件。"""
    if af.path and Path(af.path).exists():
        return str(af.path)
    if af.cos_key:
        akey = media_store.audio_track_key(af, af.session_id)
        if akey:
            url = media_store.presigned_url(akey, expires=3600)
            if url:
                return url
    part = find_part_file(af.path)
    return str(part) if part else None


def _clip_audio_by_streamer(db: Session, session_id: int, af, streamer_id: int):
    row = db.get(ReplaySession, session_id)
    turns = db.execute(
        select(StreamerTurn).where(StreamerTurn.session_id == session_id,
                                   StreamerTurn.streamer_id == streamer_id)
        .order_by(StreamerTurn.seq)
    ).scalars().all()
    if not turns:
        raise HTTPException(404, "该主播在本场没有出镜时段，无法按主播裁剪音频")
    total_ms = None
    if row and row.started_at and row.ended_at:
        total_ms = int((row.ended_at - row.started_at).total_seconds() * 1000)

    fname = _media_filename(db, row, streamer_id, "audio") if row else f"streamer{streamer_id}.m4a"
    # 该主播几乎覆盖整场（单段且全程）→ 直接返回整场，免裁剪
    if len(turns) == 1:
        t0 = turns[0]
        covers_all = (t0.start_ms or 0) <= 1500 and (
            not t0.end_ms or t0.end_ms <= t0.start_ms
            or (total_ms and t0.end_ms >= total_ms - 1500))
        if covers_all:
            if af.cos_key:
                url = media_store.presigned_url(af.cos_key, expires=7200, filename=fname)
                if url:
                    return RedirectResponse(url)
            if af.path and Path(af.path).exists():
                return FileResponse(af.path, filename=fname, media_type="audio/mp4")

    source = _audio_source(af)
    if not source:
        raise HTTPException(404, "源音频不可用（本地已清理且无 COS 备份）")

    tmpdir = tempfile.mkdtemp(prefix=f"clip_{session_id}_{streamer_id}_")
    seg_files: list[str] = []
    try:
        for i, t in enumerate(turns):
            start = max(0, t.start_ms) / 1000.0
            end_ms = t.end_ms if (t.end_ms and t.end_ms > t.start_ms) else (total_ms or 0)
            dur_args = []
            if end_ms and end_ms > t.start_ms:
                dur_args = ["-t", f"{(end_ms - t.start_ms) / 1000.0:.3f}"]
            out = os.path.join(tmpdir, f"seg{i}.m4a")
            # -c copy 远程 remux（音频 aac 逐帧独立，秒级、精度足够）
            cmd = ["ffmpeg", "-y", "-hide_banner", "-loglevel", "error",
                   "-ss", f"{start:.3f}", "-i", source, *dur_args,
                   "-vn", "-c", "copy", "-movflags", "+faststart", out]
            r = subprocess.run(cmd, capture_output=True, timeout=600)
            if r.returncode == 0 and os.path.exists(out) and os.path.getsize(out) > 0:
                seg_files.append(out)
        # 任一时段裁剪失败即视为不完整，宁可明确报错也不静默给残缺音频
        if len(seg_files) < len(turns):
            raise HTTPException(
                502, f"该主播有 {len(turns)} 个出镜时段，仅成功裁剪 {len(seg_files)} 段，"
                     f"结果不完整。请稍后重试或改导出整场音频。")
        if len(seg_files) == 1:
            final = seg_files[0]
        else:
            final = os.path.join(tmpdir, "final.m4a")
            listf = os.path.join(tmpdir, "list.txt")
            with open(listf, "w") as f:
                for s in seg_files:
                    f.write(f"file '{s}'\n")
            subprocess.run(["ffmpeg", "-y", "-hide_banner", "-loglevel", "error",
                            "-f", "concat", "-safe", "0", "-i", listf, "-c", "copy", final],
                           capture_output=True, timeout=600)
            if not (os.path.exists(final) and os.path.getsize(final) > 0):
                raise HTTPException(502, "多段音频拼接失败，请稍后重试或改导出整场音频。")
    except HTTPException:
        shutil.rmtree(tmpdir, ignore_errors=True)
        raise
    except subprocess.TimeoutExpired:
        shutil.rmtree(tmpdir, ignore_errors=True)
        raise HTTPException(504, "音频裁剪超时（源较大或网络较慢），请稍后重试或改导出整场音频。")
    except Exception:
        shutil.rmtree(tmpdir, ignore_errors=True)
        raise HTTPException(500, "音频裁剪异常")

    return FileResponse(final, filename=fname, media_type="audio/mp4",
                        background=BackgroundTask(shutil.rmtree, tmpdir, ignore_errors=True))


@router.get("/{session_id}/media_url")
def media_url(session_id: int, db: Session = Depends(get_db)):
    """给复盘工作台/分段回放用：返回可播放的媒体 URL（COS 签名或本地下载路径）+ 类型。"""
    af = get_audio(db, session_id)
    if not af or af.status in (AudioStatus.deleted, AudioStatus.error):
        return {"available": False}
    kind = getattr(af, "media_kind", "audio")
    if af.cos_key:
        if af.cos_key.endswith(".m3u8"):
            # HLS：走后端重写端点（逐切片签名），前端用 hls.js / Safari 原生播放
            return {"available": True, "kind": kind, "hls": True,
                    "url": f"/api/sessions/{session_id}/hls.m3u8"}
        url = media_store.presigned_url(af.cos_key, expires=7200)
        if url:
            return {"available": True, "kind": kind, "url": url}
    if af.path and Path(af.path).exists():
        return {"available": True, "kind": kind,
                "url": f"/api/sessions/{session_id}/audio/download"}
    return {"available": False}


@router.get("/{session_id}/hls.m3u8")
def hls_playlist(session_id: int, db: Session = Depends(get_db)):
    """HLS 播放清单：从 COS 取 index.m3u8，逐切片重写为 4h 签名 URL（私有桶方案）。"""
    af = get_audio(db, session_id)
    if (not af or af.status in (AudioStatus.deleted, AudioStatus.error)
            or not (af.cos_key or "").endswith(".m3u8")):
        raise HTTPException(404, "本场无 HLS 媒体")
    text = media_store.rewrite_hls_playlist(
        af.cos_key, expires=4 * 3600,
        per_segment_discontinuity=_legacy_reset_ts(db, session_id))
    if text is None:
        raise HTTPException(502, "云端播放清单读取失败")
    return Response(content=text, media_type="application/vnd.apple.mpegurl")


# 2026-07-07 12:00 前录的切片带 reset_timestamps(每片时间戳归零)——播放清单须
# 逐片声明 DISCONTINUITY 才能正确播；之后的新录制时间戳连续，走标准清单。
_LEGACY_TS_CUTOFF = datetime(2026, 7, 7, 12, 0, 0)


def _legacy_reset_ts(db: Session, session_id: int) -> bool:
    row = db.get(ReplaySession, session_id)
    return bool(row and row.started_at and row.started_at < _LEGACY_TS_CUTOFF)


@router.get("/{session_id}/compass-review")
def compass_review_result(session_id: int, db: Session = Depends(get_db)):
    """罗盘数据变化归因结果（收尾时算好、存成 category=罗盘归因 的整场级 Critique）。"""
    row = db.execute(select(Critique).where(
        Critique.session_id == session_id,
        Critique.level == CritiqueLevel.summary,
        Critique.category == "罗盘归因")).scalars().first()
    if not row or not row.data:
        return {"available": False}
    try:
        return {"available": True, **json.loads(row.data)}
    except json.JSONDecodeError:
        return {"available": False}


@router.delete("/{session_id}/audio")
def delete_audio(session_id: int, db: Session = Depends(get_db)):
    removed, msg = hard_delete_audio(db, session_id)
    return {"removed": removed, "message": msg}


@router.delete("/{session_id}")
def delete_session(session_id: int, db: Session = Depends(get_db)):
    """彻底删除整个场次（转写/点评/OCR/班次/录音一并删除，不可恢复）。"""
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    stop_pipeline(session_id)   # 若在跑先停
    if row.audio:
        if row.audio.path:
            Path(row.audio.path).unlink(missing_ok=True)
        if row.audio.cos_key:
            # HLS 场次要删整个切片目录前缀——只删 index.m3u8 会让几百个 .ts 切片
            # 泄漏在 COS（继续计费、隐私未清）。与 hard_delete_audio / _retention_sweep 对齐。
            if row.audio.cos_key.endswith(".m3u8"):
                media_store.delete_prefix(f"replay/session_{session_id}/")
                media_store.delete(f"replay/session_{session_id}.m4a")
            else:
                media_store.delete(row.audio.cos_key)   # 单文件模式：直接删对象
    db.delete(row)              # cascade 删除转写/点评/OCR/班次/音频记录
    db.commit()
    return {"deleted": True}


def _wt(started_at, start_ms: int) -> str:
    """相对毫秒 → 墙钟 HH:MM:SS（无开播时间则退化为 mm:ss）。"""
    if started_at:
        return (started_at + timedelta(milliseconds=start_ms)).strftime("%H:%M:%S")
    s = (start_ms or 0) // 1000
    return f"{s // 60:02d}:{s % 60:02d}"


def _rhythm_line(data: str) -> str:
    """段级 critique.data 里的节奏指标 → 一行人读文本（导出用）。"""
    try:
        m = (json.loads(data) or {}).get("metrics") if data else None
    except json.JSONDecodeError:
        m = None
    if not m:
        return ""
    cv = m.get("energy_cv")
    extra = f" · 响度起伏 {cv}" if cv is not None else ""
    return (f"语速 {m.get('wpm', 0)} 字/分 · 冷场 {m.get('cold_opens', 0)} 次 · "
            f"停顿均 {m.get('avg_pause_ms', 0)}ms · 重复 {int(m.get('repeat_ratio', 0) * 100)}% · "
            f"互动 {m.get('ask_count', 0)} · 弹幕 {m.get('danmaku', 0)}{extra}")


def _emotion_of(data: str) -> str:
    """段级 critique.data 里的语气情绪标签（优先 Qwen 识别，其次 LLM 判定）。"""
    try:
        d = (json.loads(data) or {}) if data else {}
    except json.JSONDecodeError:
        return ""
    return d.get("asr_emotion_label") or d.get("emotion") or ""


def _emotion_match_line(data: str) -> str:
    """R8 情绪-语境对齐结论（导出用）。"""
    try:
        match = ((json.loads(data) or {}).get("emotion_match") or {}) if data else {}
    except json.JSONDecodeError:
        return ""
    if not match:
        return ""
    status = "匹配" if match.get("matched") else "不匹配"
    line = (f"情绪-语境：{status} · 实际 {match.get('actual_emotion', '-')} · "
            f"应有 {match.get('expected_emotion', '-')}")
    if match.get("reason"):
        line += f" · {match['reason']}"
    if not match.get("matched") and match.get("suggestion"):
        line += f" · 建议：{match['suggestion']}"
    return line


def _scene_of(data: str) -> str:
    """段级 critique.data 里的话术场景标签。"""
    try:
        return ((json.loads(data) or {}).get("scene_label") or "") if data else ""
    except json.JSONDecodeError:
        return ""


def _pitch_line(data: str) -> str:
    """段级语调（真实音高 F0）一行。"""
    try:
        pt = (json.loads(data) or {}).get("pitch") if data else None
    except json.JSONDecodeError:
        pt = None
    if not pt or not pt.get("pitch_hz"):
        return ""
    return f"语调 {pt.get('label', '')} · 均值 {pt['pitch_hz']}Hz · 起伏 {pt.get('pitch_cv', 0)}"


def _loop_line(data: str) -> str:
    """段级讲款闭环缺失提示（一行，仅讲款段且有缺失时）。"""
    try:
        lp = (json.loads(data) or {}).get("loop") if data else None
    except json.JSONDecodeError:
        lp = None
    if not lp or not lp.get("missing"):
        return ""
    return f"讲款闭环 {lp.get('pct', 0)}% · 缺：{'、'.join(lp['missing'])}"


def _fmt_dt(dt) -> str:
    return dt.strftime("%Y-%m-%d %H:%M:%S") if dt else "-"


def _export_data(db: Session, row: ReplaySession):
    sname = {s.id: s.name for s in db.query(Streamer).all()}
    def nm(sid):
        return (sname.get(sid) if sid else None) or "未指定主播"
    trs = db.execute(select(Transcript).where(Transcript.session_id == row.id)
                     .order_by(Transcript.seq)).scalars().all()
    crs = db.execute(select(Critique).where(Critique.session_id == row.id)).scalars().all()
    ocr = db.execute(select(OcrRecord).where(OcrRecord.session_id == row.id)
                     .order_by(OcrRecord.start_ms)).scalars().all()
    segs = sorted((c for c in crs if c.level == CritiqueLevel.segment), key=lambda c: c.start_ms)
    sflags = [c for c in crs if c.level == CritiqueLevel.sentence]
    summaries = [c for c in crs
                 if c.level == CritiqueLevel.summary and c.category == "整场总结"]
    rhythm_rows = db.execute(select(SessionRhythmAnalysis).where(
        SessionRhythmAnalysis.session_id == row.id,
    )).scalars().all()
    rhythm_by_streamer = {analysis.streamer_id: analysis for analysis in rhythm_rows}
    projected = []
    covered_streamers = set()
    for summary in summaries:
        analysis = rhythm_by_streamer.get(summary.streamer_id)
        data_text = (
            _projected_summary_data(summary.data, analysis)
            if analysis is not None
            else json.dumps(_json_dict(summary.data), ensure_ascii=False)
        )
        projected.append(SimpleNamespace(
            streamer_id=summary.streamer_id,
            score=summary.score,
            comment=summary.comment,
            data=data_text,
        ))
        covered_streamers.add(summary.streamer_id)
    # 主总结偶发失败时，独立 R3 事实仍必须能被查看和导出。
    for analysis in rhythm_rows:
        if analysis.streamer_id in covered_streamers:
            continue
        projected.append(SimpleNamespace(
            streamer_id=analysis.streamer_id,
            score=None,
            comment="R3 话术节奏分析已生成",
            data=_projected_summary_data(None, analysis),
        ))
    summaries = projected
    return nm, trs, ocr, segs, sflags, summaries


def _loads(data) -> dict:
    try:
        parsed = (json.loads(data) or {}) if data else {}
    except (json.JSONDecodeError, TypeError):
        return {}
    return parsed if isinstance(parsed, dict) else {}


def _item_text(x) -> str:
    """top_issues / action_items 里既可能是字符串、也可能是 {action,source} 对象，统一取正文。"""
    if isinstance(x, dict):
        return str(x.get("action") or x.get("text") or x.get("issue")
                   or x.get("title") or x.get("content") or "")
    return str(x if x is not None else "")


def _item_src(x) -> str:
    return str(x.get("source") or "") if isinstance(x, dict) else ""


def _experience_sources(data: dict) -> list[dict]:
    rows = data.get("experience_sources")
    return [row for row in rows if isinstance(row, dict)] if isinstance(rows, list) else []
_RHYTHM_DIFF_LABELS = {
    "matched": "符合",
    "missing": "缺失",
    "too_long": "过长",
    "too_short": "过短",
    "wrong_content": "内容偏离",
    "out_of_order": "顺序不对",
}


def _as_dict(value) -> dict:
    return value if isinstance(value, dict) else {}


def _dict_list(value) -> list[dict]:
    return [item for item in value if isinstance(item, dict)] if isinstance(value, list) else []


def _rhythm_analysis_md(L, rhythm) -> None:
    """把 R3 的真实结构、逐段差异和验证结论写进 Markdown 报告。"""
    if not isinstance(rhythm, dict) or not rhythm.get("available"):
        return
    template = _as_dict(rhythm.get("template"))
    actual = _as_dict(rhythm.get("actual"))
    comparison = _as_dict(rhythm.get("comparison"))
    validation = _as_dict(rhythm.get("validation"))
    match = comparison.get("match_score")
    title = template.get("title") or "暂无可用优秀模板"
    L.append("\n**R3 话术节奏分析与优秀模板对比**")
    L.append(f"- 对比模板：《{title}》" + (f" · 匹配度 {match} 分" if match is not None else ""))
    if rhythm.get("error_msg"):
        prefix = "计划已停止归因" if rhythm.get("declared_execution") else "本轮分析未更新"
        L.append(f"- 当前状态：{prefix}；{rhythm['error_msg']}")
    if comparison.get("summary"):
        L.append(f"- 对比结论：{comparison['summary']}")
    blocks = _dict_list(actual.get("blocks"))
    if blocks:
        L.append("\n**本场实际节奏结构**")
        for block in blocks:
            duration = block.get("duration_sec", 0)
            L.append(
                f"- `{block.get('time_range', '-')}` {block.get('scene_label') or block.get('scene') or '未分类'}"
                f"（{duration} 秒）：{block.get('content') or block.get('quote') or '-'}"
            )
    differences = _dict_list(comparison.get("differences"))
    if differences:
        L.append("\n**与优秀节奏逐段差异**")
        for item in differences:
            label = _RHYTHM_DIFF_LABELS.get(item.get("status"), item.get("status") or "待判断")
            expected = item.get("expected_duration_sec")
            actual_duration = item.get("actual_duration_sec")
            timing = ""
            if expected is not None or actual_duration is not None:
                timing = f"（模板 {expected or 0} 秒 / 实际 {actual_duration or 0} 秒）"
            L.append(
                f"- 第 {item.get('step_order', '-')} 段 · {item.get('template_phase') or '模板步骤'} · {label}{timing}："
                f"{item.get('difference') or '-'}"
            )
            expected_text = item.get("expected_content") or item.get("expected_action")
            if expected_text:
                L.append(f"  - 模板要求：{expected_text}")
            evidence = _dict_list(item.get("evidence"))
            if evidence:
                L.append("  - 本场证据：" + "；".join(
                    f"`{row.get('time_range', '-')}` {row.get('quote') or '-'}" for row in evidence
                ))
            else:
                L.append("  - 本场证据：未找到对应逐字稿片段")
            if item.get("action"):
                L.append(f"  - 下一场这样做：{item['action']}")
    extras = _dict_list(comparison.get("extra_issues"))
    for item in extras:
        L.append(f"- 额外节奏问题：{item.get('diagnosis') or '-'}")
        if item.get("action"):
            L.append(f"  - 下一场这样做：{item['action']}")
    if validation:
        verdict = validation.get("verdict") or validation.get("status") or "待验证"
        L.append(f"\n**执行效果验证**：{verdict}")
        if validation.get("reason"):
            L.append(f"- {validation['reason']}")
        for evidence in _dict_list(validation.get("evidence"))[:8]:
            L.append(f"- `{evidence.get('time', '-')}` {evidence.get('effect') or evidence.get('metric') or '-'}")


def _rhythm_analysis_docx(doc, rhythm) -> None:
    """把 R3 结果写进 Word 报告，内容与 Markdown 版保持一致。"""
    if not isinstance(rhythm, dict) or not rhythm.get("available"):
        return
    template = _as_dict(rhythm.get("template"))
    actual = _as_dict(rhythm.get("actual"))
    comparison = _as_dict(rhythm.get("comparison"))
    validation = _as_dict(rhythm.get("validation"))
    doc.add_heading("R3 话术节奏分析与优秀模板对比", level=3)
    title = template.get("title") or "暂无可用优秀模板"
    match = comparison.get("match_score")
    doc.add_paragraph(f"对比模板：《{title}》" + (f" · 匹配度 {match} 分" if match is not None else ""))
    if rhythm.get("error_msg"):
        prefix = "计划已停止归因" if rhythm.get("declared_execution") else "本轮分析未更新"
        doc.add_paragraph(f"当前状态：{prefix}；{rhythm['error_msg']}")
    if comparison.get("summary"):
        doc.add_paragraph("对比结论：" + str(comparison["summary"]))
    blocks = _dict_list(actual.get("blocks"))
    if blocks:
        doc.add_heading("本场实际节奏结构", level=4)
        for block in blocks:
            doc.add_paragraph(
                f"[{block.get('time_range', '-')}] {block.get('scene_label') or block.get('scene') or '未分类'}"
                f"（{block.get('duration_sec', 0)} 秒）：{block.get('content') or block.get('quote') or '-'}",
                style="List Bullet",
            )
    differences = _dict_list(comparison.get("differences"))
    if differences:
        doc.add_heading("与优秀节奏逐段差异", level=4)
        for item in differences:
            label = _RHYTHM_DIFF_LABELS.get(item.get("status"), item.get("status") or "待判断")
            expected = item.get("expected_duration_sec")
            actual_duration = item.get("actual_duration_sec")
            timing = ""
            if expected is not None or actual_duration is not None:
                timing = f"（模板 {expected or 0} 秒 / 实际 {actual_duration or 0} 秒）"
            doc.add_paragraph(
                f"第 {item.get('step_order', '-')} 段 · {item.get('template_phase') or '模板步骤'} · {label}{timing}："
                f"{item.get('difference') or '-'}",
                style="List Bullet",
            )
            expected_text = item.get("expected_content") or item.get("expected_action")
            if expected_text:
                doc.add_paragraph("模板要求：" + str(expected_text), style="List Bullet 2")
            evidence = _dict_list(item.get("evidence"))
            if evidence:
                doc.add_paragraph("本场证据：" + "；".join(
                    f"[{row.get('time_range', '-')}] {row.get('quote') or '-'}" for row in evidence
                ), style="List Bullet 2")
            else:
                doc.add_paragraph("本场证据：未找到对应逐字稿片段", style="List Bullet 2")
            if item.get("action"):
                doc.add_paragraph("下一场这样做：" + str(item["action"]), style="List Bullet 2")
    for item in _dict_list(comparison.get("extra_issues")):
        doc.add_paragraph("额外节奏问题：" + str(item.get("diagnosis") or "-"), style="List Bullet")
        if item.get("action"):
            doc.add_paragraph("下一场这样做：" + str(item["action"]), style="List Bullet 2")
    if validation:
        doc.add_heading("执行效果验证", level=4)
        doc.add_paragraph(str(validation.get("verdict") or validation.get("status") or "待验证"))
        if validation.get("reason"):
            doc.add_paragraph(str(validation["reason"]))
        for evidence in _dict_list(validation.get("evidence"))[:8]:
            doc.add_paragraph(
                f"[{evidence.get('time', '-')}] {evidence.get('effect') or evidence.get('metric') or '-'}",
                style="List Bullet",
            )


def _order_sids(summaries, segs, trs) -> list[int]:
    seen: list[int] = []
    for coll in (summaries, segs, trs):
        for x in coll:
            if x.streamer_id and x.streamer_id not in seen:
                seen.append(x.streamer_id)
    return seen


# ---- Markdown 渲染片段 ----

def _summary_md(L, c, nm) -> None:
    L.append(f"\n### {nm(c.streamer_id)}　综合 {c.score if c.score is not None else '-'} 分")
    d = _loads(c.data)
    if d.get("dimensions"):
        L.append("- 维度评分：" + " ｜ ".join(f"{k} {v}" for k, v in d["dimensions"].items()))
    if d.get("emotion"):
        L.append(f"- 整体语气情绪：{d['emotion']}")
    ov = d.get("emotion_overview") or {}
    if ov.get("dist"):
        L.append("- 语音情绪分布(Qwen)：" + " · ".join(f"{x['label']} {x['pct']}%" for x in ov["dist"])
                 + (f"（主导：{ov['dominant']}）" if ov.get("dominant") else ""))
    st = d.get("structure") or {}
    if st.get("arc"):
        L.append(f"- 整场环节结构：{' → '.join(st['arc'])}")
        if st.get("distribution"):
            L.append("- 各环节占比：" + "；".join(f"{x['scene']} {x['pct']}%" for x in st["distribution"]))
    if d.get("structure_review"):
        L.append(f"- 话术逻辑点评：{d['structure_review']}")
    for x in (st.get("diagnostics") or []):
        L.append(f"  - ⚠️ {x}")
    lp = d.get("loop") or {}
    if lp.get("product_segments"):
        L.append(f"- 讲款闭环：讲款段 {lp['product_segments']} 个 · 平均完整度 {lp.get('avg_completeness', 0)}%"
                 + (f" · 常缺环节：{'、'.join(lp['often_missing'])}" if lp.get("often_missing") else " · 环节基本齐全"))
    _rhythm_analysis_md(L, d.get("rhythm_analysis"))
    if c.comment:
        L.append(f"\n{c.comment}")
    refs = _experience_sources(d)
    if refs:
        L.append("\n**个性化经验来源**")
        for ref in refs:
            status = "已引用" if ref.get("used") else "已召回"
            citation = ref.get("citation") or ref.get("ref") or "-"
            source = ref.get("source") or "未标明来源"
            scene = f" · 场景 {ref['scene']}" if ref.get("scene") else ""
            L.append(f"- {status} `{citation}` · {source}{scene}")
    if d.get("top_issues"):
        L.append("\n**Top 问题**")
        L += [f"{i + 1}. {_item_text(x)}" for i, x in enumerate(d["top_issues"])]
    if d.get("action_items"):
        L.append("\n**改进行动项**")
        for i, x in enumerate(d["action_items"]):
            src = _item_src(x)
            L.append(f"{i + 1}. {_item_text(x)}" + (f"（来源：{src}）" if src else ""))


def _timeline_md(L, started, trs, segs, sflags, ocr, nm) -> None:
    used_t, used_f, used_d = set(), set(), set()
    if segs:
        for i, seg in enumerate(segs):
            b = 10 ** 15 if i == len(segs) - 1 else segs[i + 1].start_ms
            L.append(f"\n#### [{_wt(started, seg.start_ms)} – {_wt(started, seg.end_ms)}]　"
                     f"段级点评 {seg.score if seg.score is not None else ''}分")
            sc = _scene_of(seg.data)
            if sc:
                L.append(f"> 话术场景：{sc}")
            ll = _loop_line(seg.data)
            if ll:
                L.append(f"> {ll}")
            if seg.comment:
                L.append(f"> {seg.comment}")
            rl = _rhythm_line(seg.data)
            if rl:
                L.append(f"> 节奏：{rl}")
            emo = _emotion_of(seg.data)
            if emo:
                L.append(f"> 语气情绪：{emo}")
            eml = _emotion_match_line(seg.data)
            if eml:
                L.append(f"> {eml}")
            pl = _pitch_line(seg.data)
            if pl:
                L.append(f"> {pl}")
            L.append("\n**话术：**")
            any_t = False
            for t in trs:
                if t.start_ms <= b and t.id not in used_t:
                    used_t.add(t.id)
                    any_t = True
                    L.append(f"- `{_wt(started, t.start_ms)}` [{nm(t.streamer_id)}] {t.corrected_text}")
            if not any_t:
                L.append("- （本段无转写；可能无人说话或识别未产出）")
            fl = [c for c in sflags if c.start_ms <= b and c.id not in used_f]
            if fl:
                L.append("\n**句级硬错：**")
                for c in fl:
                    used_f.add(c.id)
                    L.append(f"- `{_wt(started, c.start_ms)}` ⚠️ [{c.category}] {c.comment}（原话：{c.quote}）")
            dm = [o for o in ocr if o.start_ms <= b and o.id not in used_d]
            if dm:
                L.append("\n**弹幕：**")
                for o in dm:
                    used_d.add(o.id)
                    L.append(f"- `{_wt(started, o.start_ms)}` {o.text}")
    elif trs:
        L.append("")
        for t in trs:
            L.append(f"- `{_wt(started, t.start_ms)}` [{nm(t.streamer_id)}] {t.corrected_text}")
    else:
        L.append("\n_（无话术记录。）_")


def _build_markdown(db: Session, row: ReplaySession,
                    only_sid: int | None = None, by_role: bool = False) -> str:
    nm, trs, ocr, segs, sflags, summaries = _export_data(db, row)
    started = row.started_at
    note = (f"（仅主播：{nm(only_sid)}）" if only_sid
            else "（按主播角色分章）" if by_role else "")

    L = [f"# {row.title or '直播复盘'}　(场次 #{row.id}){note}", ""]
    L.append(f"- 开播：{_fmt_dt(started)}　结束：{_fmt_dt(row.ended_at)}")
    if row.room_id:
        L.append(f"- 房间号：{row.room_id}")
    L.append(f"- 统计：{len(trs)} 句话术 · {len(segs)} 个段级点评 · "
             f"{len(sflags)} 处硬错 · {len(ocr)} 条弹幕")
    try:
        _multi = len(json.loads(row.lineup or "[]")) >= 2
    except (json.JSONDecodeError, TypeError):
        _multi = False
    if _multi:   # 仅多主播场次才谈场控双线（单主播无场控概念，避免误导）
        from ..core.roles import analyze_roles
        _rl = analyze_roles(trs)
        if _rl["assist_count"] > 0:
            L.append(f"- 主播/场控分工：主讲 {_rl['host_count']} 句 · 场控搭话 "
                     f"{_rl['assist_count']} 句（{_rl['assist_ratio']}%）")

    if only_sid or by_role:
        ids = [only_sid] if only_sid else _order_sids(summaries, segs, trs)
        # by_role 是全量重排：若存在未指定主播(streamer_id=None)的内容，补一章兜底，避免静默丢句
        if by_role and any(x.streamer_id is None for x in (*trs, *segs, *sflags)):
            ids = ids + [None]
        if not ids:
            L.append("\n_（无主播归属数据可分。）_")
            return "\n".join(L)
        for sid in ids:
            L.append(f"\n## 主播：{nm(sid)}")
            csum = next((c for c in summaries if c.streamer_id == sid), None)
            if csum:
                _summary_md(L, csum, nm)
            elif sid is not None:
                L.append("\n_（该主播无整场报告）_")
            L.append(f"\n### 时间轴明细 · {nm(sid)}")
            _timeline_md(L, started,
                         [t for t in trs if t.streamer_id == sid],
                         [c for c in segs if c.streamer_id == sid],
                         [c for c in sflags if c.streamer_id == sid],
                         [o for o in ocr if o.streamer_id == sid], nm)
        return "\n".join(L)

    # 整场
    L.append("\n## 一、整场复盘报告")
    if summaries:
        for c in summaries:
            _summary_md(L, c, nm)
    else:
        L.append("\n_（本场未生成整场报告：可能未点「结束采集」正常收尾，或话术过少。）_")
    L.append("\n## 二、时间轴逐段明细")
    _timeline_md(L, started, trs, segs, sflags, ocr, nm)
    return "\n".join(L)


# ---- Word 渲染片段 ----

def _summary_docx(doc, c, nm) -> None:
    doc.add_heading(f"{nm(c.streamer_id)}　综合 {c.score if c.score is not None else '-'} 分", level=2)
    d = _loads(c.data)
    if d.get("dimensions"):
        doc.add_paragraph("维度评分：" + " ｜ ".join(f"{k} {v}" for k, v in d["dimensions"].items()))
    if d.get("emotion"):
        doc.add_paragraph("整体语气情绪：" + str(d["emotion"]))
    ov = d.get("emotion_overview") or {}
    if ov.get("dist"):
        doc.add_paragraph("语音情绪分布(Qwen)：" + " · ".join(f"{x['label']} {x['pct']}%" for x in ov["dist"])
                          + (f"（主导：{ov['dominant']}）" if ov.get("dominant") else ""))
    st = d.get("structure") or {}
    if st.get("arc"):
        doc.add_paragraph("整场环节结构：" + " → ".join(st["arc"]))
        if st.get("distribution"):
            doc.add_paragraph("各环节占比：" + "；".join(f"{x['scene']} {x['pct']}%" for x in st["distribution"]))
    if d.get("structure_review"):
        doc.add_paragraph("话术逻辑点评：" + str(d["structure_review"]))
    for x in (st.get("diagnostics") or []):
        doc.add_paragraph("⚠️ " + str(x), style="List Bullet")
    lp = d.get("loop") or {}
    if lp.get("product_segments"):
        doc.add_paragraph(
            f"讲款闭环：讲款段 {lp['product_segments']} 个 · 平均完整度 {lp.get('avg_completeness', 0)}%"
            + (f" · 常缺环节：{'、'.join(lp['often_missing'])}" if lp.get("often_missing") else " · 环节基本齐全"))
    _rhythm_analysis_docx(doc, d.get("rhythm_analysis"))
    if c.comment:
        doc.add_paragraph(c.comment)
    refs = _experience_sources(d)
    if refs:
        doc.add_heading("个性化经验来源", level=3)
        for ref in refs:
            status = "已引用" if ref.get("used") else "已召回"
            citation = ref.get("citation") or ref.get("ref") or "-"
            source = ref.get("source") or "未标明来源"
            scene = f" · 场景 {ref['scene']}" if ref.get("scene") else ""
            doc.add_paragraph(f"{status} {citation} · {source}{scene}", style="List Bullet")
    if d.get("top_issues"):
        doc.add_heading("Top 问题", level=3)
        for x in d["top_issues"]:
            doc.add_paragraph(_item_text(x), style="List Number")
    if d.get("action_items"):
        doc.add_heading("改进行动项", level=3)
        for x in d["action_items"]:
            src = _item_src(x)
            doc.add_paragraph(_item_text(x) + (f"（来源：{src}）" if src else ""), style="List Number")


def _timeline_docx(doc, started, trs, segs, sflags, ocr, nm) -> None:
    used_t, used_f, used_d = set(), set(), set()
    if segs:
        for i, seg in enumerate(segs):
            b = 10 ** 15 if i == len(segs) - 1 else segs[i + 1].start_ms
            doc.add_heading(f"[{_wt(started, seg.start_ms)} – {_wt(started, seg.end_ms)}]　"
                            f"段级点评 {seg.score if seg.score is not None else ''} 分", level=3)
            sc = _scene_of(seg.data)
            if sc:
                doc.add_paragraph("话术场景：" + sc)
            ll = _loop_line(seg.data)
            if ll:
                doc.add_paragraph(ll)
            if seg.comment:
                doc.add_paragraph(seg.comment)
            rl = _rhythm_line(seg.data)
            if rl:
                doc.add_paragraph("节奏：" + rl)
            emo = _emotion_of(seg.data)
            if emo:
                doc.add_paragraph("语气情绪：" + emo)
            eml = _emotion_match_line(seg.data)
            if eml:
                doc.add_paragraph(eml)
            pl = _pitch_line(seg.data)
            if pl:
                doc.add_paragraph(pl)
            doc.add_heading("话术", level=4)
            any_t = False
            for t in trs:
                if t.start_ms <= b and t.id not in used_t:
                    used_t.add(t.id)
                    any_t = True
                    p = doc.add_paragraph(style="List Bullet")
                    p.add_run(f"[{_wt(started, t.start_ms)}] {nm(t.streamer_id)}：").bold = True
                    p.add_run(t.corrected_text)
            if not any_t:
                doc.add_paragraph("（本段无转写；可能无人说话或识别未产出）")
            fl = [c for c in sflags if c.start_ms <= b and c.id not in used_f]
            if fl:
                doc.add_heading("句级硬错", level=4)
                for c in fl:
                    used_f.add(c.id)
                    doc.add_paragraph(f"⚠️ [{c.category}] {c.comment}（原话：{c.quote}）",
                                      style="List Bullet")
            dm = [o for o in ocr if o.start_ms <= b and o.id not in used_d]
            if dm:
                doc.add_heading("弹幕", level=4)
                for o in dm:
                    used_d.add(o.id)
                    doc.add_paragraph(f"[{_wt(started, o.start_ms)}] {o.text}", style="List Bullet")
    elif trs:
        for t in trs:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(f"[{_wt(started, t.start_ms)}] {nm(t.streamer_id)}：").bold = True
            p.add_run(t.corrected_text)
    else:
        doc.add_paragraph("（无话术记录。）")


def _build_docx(db: Session, row: ReplaySession,
                only_sid: int | None = None, by_role: bool = False) -> bytes:
    from docx import Document

    nm, trs, ocr, segs, sflags, summaries = _export_data(db, row)
    started = row.started_at
    note = (f"（仅主播：{nm(only_sid)}）" if only_sid
            else "（按主播角色分章）" if by_role else "")
    doc = Document()
    doc.add_heading((row.title or "直播复盘") + note, level=0)
    doc.add_paragraph(f"场次 #{row.id}　开播 {_fmt_dt(started)}　结束 {_fmt_dt(row.ended_at)}"
                      + (f"　房间 {row.room_id}" if row.room_id else ""))
    doc.add_paragraph(f"统计：{len(trs)} 句话术 · {len(segs)} 段点评 · "
                      f"{len(sflags)} 处硬错 · {len(ocr)} 条弹幕")

    if only_sid or by_role:
        ids = [only_sid] if only_sid else _order_sids(summaries, segs, trs)
        if by_role and any(x.streamer_id is None for x in (*trs, *segs, *sflags)):
            ids = ids + [None]
        if not ids:
            doc.add_paragraph("（无主播归属数据可分。）")
        for sid in ids:
            doc.add_heading(f"主播：{nm(sid)}", level=1)
            csum = next((c for c in summaries if c.streamer_id == sid), None)
            if csum:
                _summary_docx(doc, csum, nm)
            elif sid is not None:
                doc.add_paragraph("（该主播无整场报告）")
            doc.add_heading(f"时间轴明细 · {nm(sid)}", level=2)
            _timeline_docx(doc, started,
                           [t for t in trs if t.streamer_id == sid],
                           [c for c in segs if c.streamer_id == sid],
                           [c for c in sflags if c.streamer_id == sid],
                           [o for o in ocr if o.streamer_id == sid], nm)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    doc.add_heading("一、整场复盘报告", level=1)
    if summaries:
        for c in summaries:
            _summary_docx(doc, c, nm)
    else:
        doc.add_paragraph("（本场未生成整场报告：可能未点「结束采集」正常收尾，或话术过少。）")
    doc.add_heading("二、时间轴逐段明细", level=1)
    _timeline_docx(doc, started, trs, segs, sflags, ocr, nm)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


_ILLEGAL_FN = re.compile(r'[\\/:*?"<>|\n\r\t]+')


def _safe_name(s: str, limit: int = 40) -> str:
    s = _ILLEGAL_FN.sub("", (s or "").strip()).replace(" ", "")
    return s[:limit]


def _export_basename(row: ReplaySession) -> str:
    """『直播间名_MMDD_HHMM』——文件名前缀，一眼看出是哪个直播间哪场。"""
    title = (row.title or "").split(" · ")[0].strip()      # 去掉标题里的 "· 时间 复盘" 后缀
    base = _safe_name(title) or (f"房间{row.room_id}" if row.room_id else "直播复盘")
    dt = row.started_at or row.created_at
    return f"{base}_{dt:%m%d_%H%M}" if dt else base


def _streamer_name(db: Session, sid: int) -> str:
    st = db.get(Streamer, sid)
    return _safe_name(st.name, 20) if st and st.name else str(sid)


def _doc_filename(db: Session, row: ReplaySession, streamer_id: int | None,
                  by_role: bool, ext: str) -> str:
    if streamer_id:
        scope = f"{_streamer_name(db, streamer_id)}复盘"
    elif by_role:
        scope = "分主播复盘"
    else:
        scope = "整场复盘"
    return f"{_export_basename(row)}_{scope}.{ext}"


def _media_filename(db: Session, row: ReplaySession, streamer_id: int | None,
                    kind: str = "audio") -> str:
    label = "直播回看" if kind == "video" else "录音"
    ext = "mp4" if kind == "video" else "m4a"
    who = f"{_streamer_name(db, streamer_id)}" if streamer_id else ""
    return f"{_export_basename(row)}_{who}{label}.{ext}"


def _cd_header(filename: str) -> str:
    """Content-Disposition：现代浏览器用 filename*(UTF-8)，老浏览器用 ASCII 兜底。"""
    ascii_fb = _ILLEGAL_FN.sub("", filename.encode("ascii", "ignore").decode()).strip() or "replay"
    return f"attachment; filename=\"{ascii_fb}\"; filename*=UTF-8''{quote(filename)}"


def _assert_streamer_in_session(db: Session, session_id: int, sid: int | None) -> None:
    """单主播导出/裁剪校验：该主播在本场须有归属数据，否则 404（避免产出误导性空文档）。"""
    if not sid:
        return
    has = (db.query(Transcript).filter(Transcript.session_id == session_id,
                                       Transcript.streamer_id == sid).first()
           or db.query(Critique).filter(Critique.session_id == session_id,
                                        Critique.streamer_id == sid).first()
           or db.query(StreamerTurn).filter(StreamerTurn.session_id == session_id,
                                            StreamerTurn.streamer_id == sid).first()
           or db.query(SessionRhythmAnalysis).filter(
               SessionRhythmAnalysis.session_id == session_id,
               SessionRhythmAnalysis.streamer_id == sid,
           ).first())
    if not has:
        raise HTTPException(404, "该主播在本场无复盘数据")


@router.get("/{session_id}/export.md")
def export_markdown(session_id: int, streamer_id: int | None = None,
                    by_role: bool = False, db: Session = Depends(get_db)):
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    _assert_streamer_in_session(db, session_id, streamer_id)
    md = _build_markdown(db, row, only_sid=streamer_id, by_role=by_role)
    fname = _doc_filename(db, row, streamer_id, by_role, "md")
    return PlainTextResponse(md, media_type="text/markdown; charset=utf-8",
                             headers={"Content-Disposition": _cd_header(fname)})


@router.get("/{session_id}/export.docx")
def export_docx(session_id: int, streamer_id: int | None = None,
                by_role: bool = False, db: Session = Depends(get_db)):
    row = db.get(ReplaySession, session_id)
    if not row:
        raise HTTPException(404, "场次不存在")
    _assert_streamer_in_session(db, session_id, streamer_id)
    try:
        data = _build_docx(db, row, only_sid=streamer_id, by_role=by_role)
    except ModuleNotFoundError:
        raise HTTPException(500, "服务器未安装 python-docx，无法导出 Word")
    fname = _doc_filename(db, row, streamer_id, by_role, "docx")
    return Response(
        content=data,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": _cd_header(fname)},
    )


@router.post("/{session_id}/danmaku/reply")
async def danmaku_reply(session_id: int, question: str = Body(..., embed=True),
                        db: Session = Depends(get_db)):
    """场控助手：主播答不上来时，按知识库优先生成可照说的建议回答。"""
    session = db.get(ReplaySession, session_id)
    if not session:
        raise HTTPException(404, "场次不存在")
    if session.session_type == SessionType.observation:
        raise HTTPException(409, "观摩学习场次不调用实时场控大模型")
    # 场控答疑优先回答产品事实；外部优秀节奏是结构参考，不能冒充参数/价格依据。
    kb = await retrieve_detailed(
        question, top_k=5, exclude_categories=("优秀节奏", "本店经验")
    )
    kb_text = "\n".join(f"- {k['text']}" for k in kb) or "（无）"
    recent = db.execute(select(Transcript).where(Transcript.session_id == session_id)
                        .order_by(Transcript.seq.desc()).limit(4)).scalars().all()
    ctx = " ".join(t.corrected_text for t in reversed(recent))[:300] or "（无）"
    system = ("你是耳机直播间的场控助手。观众在弹幕里提问、主播一时答不上来，"
              "请给主播一句可直接照说的简洁口语回答。优先依据知识库事实，"
              "知识库没有的再据耳机常识补充，严禁编造型号/参数/价格。1-2 句、口语化。")
    user = f"【知识库】\n{kb_text}\n\n【当前讲解】{ctx}\n\n【观众提问】{question}"
    ans = await get_llm("fast").complete(system, user)
    return {"answer": (ans or "").strip()}


@router.put("/{session_id}/bind-streamer")
def bind_session_streamer(session_id: int, streamer_id: int = Body(..., embed=True),
                          force: bool = Body(False, embed=True), db: Session = Depends(get_db)):
    """R1 修复入口：绑定场次，并补齐逐字稿、OCR、点评的空归属。"""
    row = db.get(ReplaySession, session_id)
    streamer = db.get(Streamer, streamer_id)
    if row is None or streamer is None or streamer.deleted_at is not None:
        raise HTTPException(404, "场次或主播不存在")
    if row.session_type == SessionType.observation:
        raise HTTPException(409, detail={
            "code": "OBSERVATION_STREAMER_BIND_FORBIDDEN",
            "message": "观摩学习场次不绑定自家主播",
        })
    if row.status == SessionStatus.live and row.current_streamer_id not in (None, streamer_id):
        raise HTTPException(409, "直播中请使用切换主播接口")
    for model in (Transcript, OcrRecord, Critique):
        q = db.query(model).filter(model.session_id == session_id)
        if not force:
            q = q.filter(model.streamer_id.is_(None))
        q.update({model.streamer_id: streamer_id}, synchronize_session=False)
    turns = db.execute(select(StreamerTurn).where(
        StreamerTurn.session_id == session_id)).scalars().all()
    if not any(x.streamer_id == streamer_id for x in turns):
        artifacts = []
        for model in (Transcript, OcrRecord, Critique):
            artifacts.extend(db.execute(select(model).where(
                model.session_id == session_id, model.streamer_id == streamer_id)).scalars().all())
        start_ms = min([x.start_ms for x in artifacts] or [0])
        end_ms = max([x.end_ms for x in artifacts] or [0])
        db.add(StreamerTurn(session_id=session_id, streamer_id=streamer_id,
                            seq=max([x.seq for x in turns] or [-1]) + 1,
                            start_ms=start_ms, end_ms=end_ms))
    if force or not row.lineup:
        row.lineup = json.dumps([streamer_id])
    row.current_streamer_id = streamer_id if row.status == SessionStatus.live else None
    db.commit()
    return {"ok": True, "session_id": session_id, "streamer_id": streamer_id}


@router.get("/{session_id}/binding-status")
def binding_status(session_id: int, db: Session = Depends(get_db)):
    """R1 binding audit across every replay artifact."""
    row = db.get(ReplaySession, session_id)
    if row is None:
        raise HTTPException(404, "场次不存在")
    counts = {}
    owners: set[int] = set()
    for model, name in ((Transcript, "transcripts"), (OcrRecord, "ocr"), (Critique, "critiques")):
        counts[name] = db.query(model).filter(
            model.session_id == session_id, model.streamer_id.is_(None)).count()
        owners.update(x[0] for x in db.execute(select(model.streamer_id).where(
            model.session_id == session_id, model.streamer_id.is_not(None)).distinct()).all())
    turn_owners = set(x[0] for x in db.execute(select(StreamerTurn.streamer_id).where(
        StreamerTurn.session_id == session_id).distinct()).all())
    orphan_owners = sorted(owners - turn_owners)
    return {"session_id": session_id, "complete": not any(counts.values()) and not orphan_owners
            and bool(turn_owners), "unassigned": counts, "turn_streamer_ids": sorted(turn_owners),
            "orphan_streamer_ids": orphan_owners, "total_unassigned": sum(counts.values())}


@router.websocket("/{session_id}/ws")
async def session_ws(websocket: WebSocket, session_id: int):
    # WebSocket 不经过 HTTP 中间件，在这里认统一登录（或本机应急会话）
    from ..sso import authenticate
    if (await authenticate(websocket)).status != "ok":
        await websocket.close(code=1008)   # 未登录 / 没开通
        return
    await websocket.accept()
    q = bus.subscribe(session_id)
    try:
        while True:
            event = await q.get()
            await websocket.send_json(event)
    except WebSocketDisconnect:
        pass
    finally:
        bus.unsubscribe(session_id, q)
